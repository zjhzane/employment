import math
import os
import traceback
from datetime import datetime
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

URL = "https://uel-app-portal.uel.ac.uk/urd/sits.urd/run/siw_ipp_lgn.login?process=siw_ipp_app&code2=0005&code1=CD2299FAD"

email = "songshuhan@rocky2028.ccwu.cc"

password = "12345678Aa!"

_DEFAULT_XLSX = Path.home() / "Desktop" / "学生基础信息_已翻译.xlsx"

_LOGIN_BTN = (
    "#logInFormContainer > fieldset > div.sv-row > "
    "div.sv-col-sm-6.sv-col-sm-push-6 > div > input"
)

_LOGIN_TAB_LINK = (
    "#TAB-MAIN-AS > div.sv-form-container > div > div:nth-child(3) > div > "
    "ul > li:nth-child(1) > p > a"
)

# Personal information → Personal Details 标签页（性别等字段在此页）
_PERSONAL_DETAILS_TAB = "a[onclick*=\"ul_click_tab('PD')\"]"

# 下拉固定选项
_COUNTRY_OPTION_VALUE = "631"
_DISABILITY_OPTION_VALUE = "A"  # No disability
_ETHNICITY_OPTION_VALUE = "34"  # Asian - Chinese

_CONTACT_DETAILS_TAB = "#ui-id-3"
_AGENT_DETAILS_TAB = "#ui-id-4"
_QUALIFICATION_TAB = "#ui-id-6"
_EMPLOYMENT_TAB = "#ui-id-7"
_PERSONAL_STATEMENT_TAB = "#ui-id-8"
_VISA_TAB = "#ui-id-10"
_NEXT_FORM_TAB = "#ui-id-12"
_SUBMISSION_TAB = "#ui-id-13"
_FUNDING_SELF_OR_FAMILY = "SELF_OR_FAMILY"
_VISA_YES = "Y"

_SUPDOCS_UPLOAD_BTN = "#supDocs > fieldset > div > div > button"
_SUPDOCS_UPLOAD_CLOSE_BTN = "button.ui-button.ui-corner-all.ui-widget"
_DEFAULT_TRANSCRIPT_PDF = Path.home() / "Desktop" / "transcript.pdf"
_DEFAULT_CERTIFICATE_PDF = Path.home() / "Desktop" / "certficate.pdf"

# 桌面 Personal Statement：优先 PS.md → PS.docx → PS.txt（可用环境变量 UEL_PS_TXT 指定路径）
_PS_DESKTOP_CANDIDATES = ("PS.md", "PS.docx", "PS.txt")

_INSTITUTION = "Zhejiang Wanli University"
_QUALIFICATION_OFDG = "OFDG"
_MAJOR_ART_DS = "Visual Communication Design (Sino-German 2+2 double degree class)"
_SUBJECT_ART_DS = "ART-DS"
_SUBJECT_OTHER = "OTHER"
_QCOM_NO = "P"
_ENGLANG_NO = "N"
_WORK_NO = "N"
_COMPLETION_DAY = "30"
_COMPLETION_MONTH = "Jun"


def completion_year_value() -> str:
    """Expected Completion 年份：按系统当前日历年（2026→2026，进入 2027 后→2027）。"""
    return str(datetime.now().year)


def resolve_ps_path() -> Path:
    """解析 PS 文件路径：环境变量 > 桌面 PS.md / PS.docx / PS.txt（跳过 0 字节空文件）。"""
    env = os.getenv("UEL_PS_TXT", "").strip()
    if env:
        path = Path(env).expanduser()
        if path.is_file() and path.stat().st_size > 0:
            return path
        if path.is_file():
            raise ValueError(f"UEL_PS_TXT 指向的文件为空（0 字节）: {path}")
        raise FileNotFoundError(f"UEL_PS_TXT 指向的文件不存在: {path}")

    desktop = Path.home() / "Desktop"
    empty_files: list[Path] = []
    for name in _PS_DESKTOP_CANDIDATES:
        path = desktop / name
        if not path.is_file():
            continue
        if path.stat().st_size > 0:
            return path
        empty_files.append(path)

    if empty_files:
        names = ", ".join(p.name for p in empty_files)
        raise ValueError(
            f"桌面上有这些 PS 文件但内容为空（0 字节）: {names}。"
            f"请写入 Personal Statement 正文后保存，或改用 PS.md / PS.docx。"
        )
    raise FileNotFoundError(
        f"桌面上未找到 Personal Statement 文件，请放置其一: "
        f"{', '.join(_PS_DESKTOP_CANDIDATES)}"
    )


def load_ps_text(ps_path: Path) -> str:
    """读取 PS 全文（支持 .txt / .md / .docx），供 #IPQ_OA_PS 使用。"""
    path = ps_path.expanduser()
    if not path.is_file():
        raise FileNotFoundError(f"找不到 Personal Statement 文件: {path}")

    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = path.read_text(encoding="utf-8")
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as e:
            raise ImportError(
                "读取 Word（.docx）需先安装: pip install python-docx"
            ) from e
        text = "\n".join(p.text for p in Document(path).paragraphs)
    else:
        raise ValueError(
            f"不支持的 PS 格式 {suffix!r}，请使用 .txt、.md 或 .docx: {path}"
        )

    text = text.strip()
    if not text:
        size = path.stat().st_size
        raise ValueError(
            f"Personal Statement 文件无有效文字（文件大小 {size} 字节）: {path}\n"
            f"请用记事本/Word 打开该文件，粘贴 PS 全文并保存后再运行。"
        )
    return text


# 与 G 列同一数据行：AF=32, AG=33, AI=35（iter_rows min_col=7 时的下标）
_IDX_AF = 32 - 7
_IDX_AG = 33 - 7
_IDX_AI = 35 - 7


def _excel_cell_str(raw: object) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    return "" if s.lower() == "nan" else s


def _phone_from_cell(raw: object) -> str:
    if raw is None:
        return ""
    if isinstance(raw, float):
        if math.isnan(raw):
            return ""
        if raw.is_integer():
            return str(int(raw))
        return str(raw).strip()
    if isinstance(raw, int) and not isinstance(raw, bool):
        return str(raw)
    return _excel_cell_str(raw)


def _gender_code_from_cell(raw: object) -> str:
    """将 J 列性别文本映射为下拉 value：M（Male）或 F（Female）。"""
    if raw is None:
        raise ValueError("性别为空")
    s = str(raw).strip()
    if not s or s.lower() == "nan":
        raise ValueError("性别为空")
    key = s.lower().replace(" ", "")
    if key in {"f", "female", "woman", "女", "女性"} or key.startswith("female"):
        return "F"
    if key in {"m", "male", "man", "男", "男性"} or key.startswith("male"):
        return "M"
    raise ValueError(f"无法识别性别 {raw!r}，请填写 Male/Female、M/F 或 男/女")


def gender_from_xlsx(xlsx_path: Path) -> str:
    """与注册脚本一致：自第 2 行起，在 G 列有姓名的同一行读取 J 列性别。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=10, values_only=True):
            g_cell, _h, _i, j_cell = row[0], row[1], row[2], row[3]
            if g_cell is None:
                continue
            if not str(g_cell).strip():
                continue
            return _gender_code_from_cell(j_cell)
        raise ValueError(f"{xlsx_path} 的 G/J 列（自第 2 行起）没有可用的性别")
    finally:
        wb.close()


def contact_from_xlsx(xlsx_path: Path) -> tuple[str, str, str, str]:
    """与 G 列同一行：AF 按逗号分段，第 2 段为 town，第 3 段之后拼成 address；AG 邮编；AI 电话。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=35, values_only=True):
            g_cell = row[0]
            if g_cell is None:
                continue
            if not str(g_cell).strip():
                continue
            af_raw = _excel_cell_str(row[_IDX_AF])
            postcode = _excel_cell_str(row[_IDX_AG])
            phone = _phone_from_cell(row[_IDX_AI])
            if not af_raw:
                raise ValueError(f"{xlsx_path} 的 AF 列（地址）为空")
            parts = [p.strip() for p in af_raw.split(",") if p.strip()]
            town = parts[1] if len(parts) > 1 else ""
            # Address：只要第三个逗号分段之后的内容（前面三段不要）
            address = ", ".join(parts[3:]) if len(parts) > 3 else ""
            if not address:
                raise ValueError(
                    f"{xlsx_path} 的 AF 列在第三个逗号分段后没有地址内容: {af_raw!r}"
                )
            return address, postcode, town, phone
        raise ValueError(f"{xlsx_path} 的 G/AF 列（自第 2 行起）没有可用的联系信息")
    finally:
        wb.close()


def subject_from_xlsx(xlsx_path: Path) -> tuple[str, str]:
    """与 G 列同一行读取 B 列 MAJOR，返回 (下拉 value, B 列专业全文)。

    Visual Communication Design (Sino-German 2+2 double degree class) → ART-DS；
    其他专业 → OTHER，且需将 B 列内容填入 #IPQ_OA_QESUN1。
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=7, values_only=True):
            major = _excel_cell_str(row[0])
            g_cell = row[5]
            if g_cell is None:
                continue
            if not str(g_cell).strip():
                continue
            if not major:
                raise ValueError(f"{xlsx_path} 的 B 列（MAJOR）为空")
            if major == _MAJOR_ART_DS:
                return _SUBJECT_ART_DS, major
            return _SUBJECT_OTHER, major
        raise ValueError(f"{xlsx_path} 的 B/G 列（自第 2 行起）没有可用的 MAJOR")
    finally:
        wb.close()


def _open_personal_details_tab(page) -> None:
    """打开 Personal Details。链接常在 DOM 里但被隐藏，故优先调用 ul_click_tab('PD')。"""
    link = page.locator(_PERSONAL_DETAILS_TAB)
    link.wait_for(state="attached")
    switched = page.evaluate(
        """() => {
            if (typeof ul_click_tab === 'function') {
                ul_click_tab('PD');
                return true;
            }
            return false;
        }"""
    )
    if not switched:
        link.click(force=True)
    page.locator("#IPR_GEND").wait_for(state="visible")


def _select_option(page, selector: str, value: str = _COUNTRY_OPTION_VALUE) -> None:
    field = page.locator(selector)
    field.wait_for(state="visible")
    field.click()
    field.select_option(value=value)


def _fill_text(page, selector: str, text: str) -> None:
    field = page.locator(selector)
    field.wait_for(state="visible")
    field.fill(text)


def _fill_personal_statement_page(page, ps_path: Path) -> None:
    """Personal Statement 页：同面板内第一个下拉选 No，再将 PS.txt 填入 #IPQ_OA_PS。"""
    tab = page.locator(_PERSONAL_STATEMENT_TAB)
    tab.wait_for(state="visible")
    tab.click()

    ps_box = page.locator("#IPQ_OA_PS")
    ps_box.wait_for(state="visible")

    panel = page.get_by_role("tabpanel").filter(has=ps_box)
    decl = panel.locator("select").first
    decl.wait_for(state="visible")
    decl.click()
    decl.select_option(value=_WORK_NO)

    ps_box.fill(load_ps_text(ps_path))


def _visible_browse_my_computer(page):
    """定位当前可见的 Browse My Computer（id 前缀 PLUP_pickfiles，UUID 每次变）。"""
    by_text = page.get_by_role("link", name="Browse My Computer")
    for i in range(by_text.count() - 1, -1, -1):
        link = by_text.nth(i)
        if link.is_visible():
            return link

    by_id = page.locator('a[id^="PLUP_pickfiles"]')
    for i in range(by_id.count() - 1, -1, -1):
        link = by_id.nth(i)
        if link.is_visible():
            return link

    raise RuntimeError(
        "未找到可见的 Browse My Computer。请先手动点 #supDocs 上传按钮，"
        "确认展开后出现 Browse 链接。"
    )


def _click_plupload_upload_when_awaiting(page) -> None:
    """选完文件后出现 Awaiting Upload 时，点击 PLUP_uploadfiles* 的 Upload 按钮。"""
    awaiting = page.locator("p.plupupstat", has_text="Awaiting Upload")
    awaiting.first.wait_for(state="visible", timeout=30_000)

    upload_action = page.locator('a[id^="PLUP_uploadfiles"]')
    for _ in range(40):
        for i in range(upload_action.count() - 1, -1, -1):
            btn = upload_action.nth(i)
            if not btn.is_visible():
                continue
            cls = btn.get_attribute("class") or ""
            if "sv-disabled" in cls or btn.get_attribute("aria-disabled") == "true":
                continue
            btn.click()
            print("-> 已点击 Upload（PLUP_uploadfiles）开始上传")
            return
        page.wait_for_timeout(500)

    raise RuntimeError("已出现 Awaiting Upload，但 Upload 按钮一直不可点")


def _wait_plupload_success(page) -> None:
    """等待上传完成（页面出现 Successfully Uploaded）。"""
    success = page.locator("p.plupupstat", has_text="Successfully Uploaded")
    success.last.wait_for(state="visible", timeout=120_000)
    print("-> 检测到 Successfully Uploaded，可进行下一次上传")


def _close_supdocs_upload_dialog(page) -> None:
    """关闭 Supporting documents 上传弹窗（jQuery UI Close）。"""
    close_btn = page.locator(_SUPDOCS_UPLOAD_CLOSE_BTN, has_text="Close")
    for i in range(close_btn.count() - 1, -1, -1):
        btn = close_btn.nth(i)
        if btn.is_visible():
            btn.click()
            print("-> 已点击 Close 关闭上传窗口")
            page.locator(_SUPDOCS_UPLOAD_BTN).wait_for(state="visible", timeout=15_000)
            return
    raise RuntimeError("未找到可见的 Close 按钮，无法关闭上传窗口")


def _upload_supdocs_pdf(page, pdf_path: Path, label: str) -> None:
    """Supporting documents：点上传 → Browse → 选 PDF → Awaiting Upload 后点 Upload。"""
    path = pdf_path.expanduser()
    if not path.is_file():
        raise FileNotFoundError(f"找不到{label} PDF: {path}")
    resolved = str(path.resolve())

    upload_btn = page.locator(_SUPDOCS_UPLOAD_BTN)
    upload_btn.wait_for(state="visible")
    upload_btn.click()

    browse = _visible_browse_my_computer(page)
    browse.wait_for(state="visible")

    with page.expect_file_chooser(timeout=20_000) as fc_info:
        browse.click()
    fc_info.value.set_files(resolved)
    print(f"-> {label} 已选择，等待 Awaiting Upload…")

    _click_plupload_upload_when_awaiting(page)


def _fill_funding_page(page) -> None:
    """Funding 页：资金来源选 Self or Family。"""
    _select_option(page, "#IPQ_OA_TFINSUP1", value=_FUNDING_SELF_OR_FAMILY)


def _fill_submission_page(page) -> None:
    """Submission 页：Terms、Fair Processing 选 Yes，勾选同意。"""
    _select_option(page, "#IPQ_OA_TANDC", value=_VISA_YES)
    _select_option(page, "#IPQ_OA_FAIR", value=_VISA_YES)
    agree = page.locator("#IPQ_OA_AGREE")
    agree.wait_for(state="visible")
    agree.check()
    print("-> Submission 页已填写 Terms / Fair Processing / 同意勾选")


def _fill_visa_page(page) -> None:
    """Visa / immigration 页（#ui-id-10）。"""
    tab = page.locator(_VISA_TAB)
    tab.wait_for(state="visible")
    tab.click()

    _select_option(page, "#IPQ_OA_VISA", value=_VISA_YES)
    _select_option(page, "#IPQ_OA_PRV_VIS", value=_WORK_NO)
    _select_option(page, "#IPQ_OA_UKSTUD", value=_WORK_NO)
    _select_option(page, "#IPQ_OA_REFUSE", value=_WORK_NO)
    _select_option(page, "#IPQ_OA_PPT", value=_WORK_NO)


def main() -> None:
    headless = os.getenv("HEADLESS", "").strip() in {"1", "true", "True", "yes", "YES"}

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=headless)
        context = browser.new_context()
        page = context.new_page()
        run_error: BaseException | None = None
        try:
            _run_fill_form(page)
        except BaseException as e:
            run_error = e
            print("\n[错误] 脚本执行失败，浏览器将保持打开以便检查页面：")
            traceback.print_exc()
        finally:
            input(">>> 按 Enter 关闭浏览器...")
            context.close()
            browser.close()

        if run_error is not None:
            raise run_error


def _run_fill_form(page) -> None:
    page.goto(URL, wait_until="domcontentloaded")

    mua = page.locator("#MUA_CODE")
    mua.wait_for(state="visible")
    mua.fill(email)

    pw = page.locator("#PASSWORD")
    pw.wait_for(state="visible")
    pw.fill(password)

    login_btn = page.locator(_LOGIN_BTN)
    login_btn.wait_for(state="visible")
    login_btn.click()

    login_tab = page.locator(_LOGIN_TAB_LINK)
    login_tab.wait_for(state="visible")
    login_tab.click()

    _open_personal_details_tab(page)

    xlsx = Path(os.environ.get("STUDENT_INFO_XLSX", str(_DEFAULT_XLSX))).expanduser()
    gender_value = gender_from_xlsx(xlsx)

    gender = page.locator("#IPR_GEND")
    gender.wait_for(state="visible")
    gender.click()
    gender.select_option(value=gender_value)

    _select_option(page, "#IPQ_OA_COB")  # Country of birth → China

    _select_option(page, "#IPQ_OA_NAT")  # Nationality → Chinese（value 同为 631）

    _select_option(page, "#IPQ_OA_COD")  # Country of domicile → China

    _select_option(page, "#IPQ_OA_DISAB", value=_DISABILITY_OPTION_VALUE)

    _select_option(page, "#IPQ_OA_ETH", value=_ETHNICITY_OPTION_VALUE)

    contact_tab = page.locator(_CONTACT_DETAILS_TAB)
    contact_tab.wait_for(state="visible")
    contact_tab.click()

    _select_option(page, "#IPR_CODC")  # Contact Details → Country → China

    address, postcode, town, phone = contact_from_xlsx(xlsx)
    _fill_text(page, "#IPR_HAD1", address)
    _fill_text(page, "#IPR_HAPC", postcode)
    _fill_text(page, "#IPR_HAD3", town)
    _fill_text(page, "#IPR_HTEL", phone)

    agent_tab = page.locator(_AGENT_DETAILS_TAB)
    agent_tab.wait_for(state="visible")
    agent_tab.click()

    qualification_tab = page.locator(_QUALIFICATION_TAB)
    qualification_tab.wait_for(state="visible")
    qualification_tab.click()

    _fill_text(page, "#IPQ_OA_QUINS1", _INSTITUTION)
    _select_option(page, "#IPQ_OA_QEQEC1", value=_QUALIFICATION_OFDG)
    subject_value, major_text = subject_from_xlsx(xlsx)
    _select_option(page, "#IPQ_OA_QESUC1", value=subject_value)
    if subject_value == _SUBJECT_OTHER:
        other_subject = page.locator("#IPQ_OA_QESUN1")
        other_subject.wait_for(state="visible")
        _fill_text(page, "#IPQ_OA_QESUN1", major_text)

    _select_option(page, "#IPQ_OA_QCOM1", value=_QCOM_NO)

    _select_option(page, "#IPQ_OA_QENDD1_DAY", value=_COMPLETION_DAY)
    _select_option(page, "#IPQ_OA_QENDD1_MONTH", value=_COMPLETION_MONTH)
    _select_option(page, "#IPQ_OA_QENDD1_YEAR", value=completion_year_value())
    _select_option(page, "#IPQ_OA_ENGLANGT", value=_ENGLANG_NO)

    employment_tab = page.locator(_EMPLOYMENT_TAB)
    employment_tab.wait_for(state="visible")
    employment_tab.click()

    _select_option(page, "#IPQ_OA_WORK", value=_WORK_NO)

    _fill_personal_statement_page(page, resolve_ps_path())

    _fill_visa_page(page)

    transcript = Path(
        os.environ.get("UEL_TRANSCRIPT_PDF", str(_DEFAULT_TRANSCRIPT_PDF))
    ).expanduser()
    certificate = Path(
        os.environ.get("UEL_CERTIFICATE_PDF", str(_DEFAULT_CERTIFICATE_PDF))
    ).expanduser()

    _upload_supdocs_pdf(page, transcript, "成绩单 transcript")
    _wait_plupload_success(page)
    _close_supdocs_upload_dialog(page)
    _upload_supdocs_pdf(page, certificate, "证书 certificate")
    _wait_plupload_success(page)
    _close_supdocs_upload_dialog(page)

    next_tab = page.locator(_NEXT_FORM_TAB)
    next_tab.wait_for(state="visible")
    next_tab.click()
    print("-> 已点击 #ui-id-12，进入 Funding 页面")

    _fill_funding_page(page)

    submission_tab = page.locator(_SUBMISSION_TAB)
    submission_tab.wait_for(state="visible")
    submission_tab.click()
    print("-> 已点击 #ui-id-13，进入 Submission 页面")

    _fill_submission_page(page)


if __name__ == "__main__":
    main()
