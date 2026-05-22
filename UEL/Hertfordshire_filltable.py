"""赫特福德大学在线申请：打开入口、点击申请并填写个人基本信息。

数据来源：桌面「学生基础信息_已翻译.xlsx」（G/I/J 列个人信息，AF/AG/AI 列联系地址与手机）。

网申共 6 页（页眉显示 Page 1 / 6 … Page 6 / 6）；当前自动填写第 1–6 页（第 6 页勾选条款同意）。
"""

from __future__ import annotations

import math
import os
import re
import traceback
from datetime import datetime
from pathlib import Path

import openpyxl
from playwright.sync_api import FrameLocator, Page, expect, sync_playwright
from playwright.sync_api import Error as PlaywrightError

from playwright_open_uel import _dob_yyyymmdd_from_chinese_id

URL = "https://www.herts.ac.uk/international/apply/make-an-application"

email = "songshuhan@rocky2028.ccwu.cc"

_APPLY_BTN = (
    "#main > div:nth-child(11) > div > div > a > div > "
    "div._gecko-btn-cta-card-primary-body > button"
)
# 表单在 Gecko 网申 iframe 内（勿用裸 "iframe"，页上还有 Cookie 用的隐藏 iframe）
_FORM_IFRAME = 'iframe[title="Form"]'
_FORM_IFRAME_FALLBACK = 'iframe[src*="geckoform"]'
# Chosen 下拉（aria-controls 对应各字段容器 id）
_TITLE_COMBO = "form_3334-combobox-container"
_GENDER_COMBO = "form_1545-combobox-container"
_AGE_18_COMBO = "form_2194-combobox-container"
_COUNTRY_COMBO = "form_2400-combobox-container"
_NATIONALITY_COMBO = "form_2402-combobox-container"
_BIRTH_COUNTRY_COMBO = "form_17173-combobox-container"
_CRIMINAL_COMBO = "form_2357-combobox-container"
_DISABILITY_COMBO = "form_2358-combobox-container"
_FORM_2362_COMBO = "form_2362-combobox-container"
_FIRST_NAME = "#form_1543_first_name"
_LAST_NAME = "#form_1543_last_name"
_EMAIL = "#form_1607"
_EMAIL_CONFIRM = "#field1607 > div:nth-child(2) > div > div > input:nth-child(2)"
_DOB_MONTH = "#form_2409_month"
_DOB_DAY = "#form_2409_day"
_DOB_YEAR = "#form_2409_year"
_PHONE = "#form_1551"
_PHONE_FLAG = "#field1551 .selected-flag"
_PHONE_FLAG_FALLBACK = "#field1551 > div:nth-child(2) > div > div > div > div"
_PHONE_LISTBOX = "#form_1551_listbox"
_STREET = "#form_2964"
_CITY = "#form_2965"
_PROVINCE = "#form_2966"
_POSTCODE = "#form_2967"
_FORM_NEXT = "#form-next"
_FORM_NEXT_XPATH = "xpath=//*[@id='form-next']"

_INSTITUTION = "Zhejiang Wanli University"

# 第 3 页
_FORM_3683 = "#form_3683"
_FORM_2399_CONTAINER = "form_2399_container"
_FORM_2411 = "#form_2411"
_FORM_2412 = "#form_2412"
_FORM_2413 = "#form_2413"
_FORM_2414 = "#form_2414"
_FORM_3173_CONTAINER = "form_3173_container"
_TABLE_ROW = "#field2474 repeatable-table-field tbody tr"

# 第 5 页（上传）：fa-upload 开弹窗 → #inputfile_3144 选文件 → #file-upload_3144
_FORM_2582_CONTAINER = "form_2582_container"
_FIELD3144_TOGGLE_UPLOAD = (
    "#field3144 file-picker-field button.gecko-btn[ng-click*='toggleModal']"
)
_FIELD3144_INPUT_FILE = "#inputfile_3144"
_FILE_UPLOAD_CONFIRM = "#file-upload_3144"
_DEFAULT_TRANSCRIPT = Path.home() / "Desktop" / "transcript.pdf"

# 第 6 页（条款同意）
_FORM_1602_CONTAINER = "form_1602_container"
_TERMS_AGREE_OPTION = "Yes, I agree to the above terms"
_DEFAULT_CERTIFICATE = Path.home() / "Desktop" / "certficate.pdf"

_ESSAY_2411 = (
    "I am eager to apply for this course at the University of Hertfordshire "
    "because of its excellent reputation for combining academic rigor with "
    "practical, industry-focused learning. Hertfordshire's strong links with "
    "leading enterprises and its emphasis on employability perfectly align with "
    "my career goals. This program offers a comprehensive curriculum that "
    "addresses current industry challenges, which will help me bridge the gap "
    "between theoretical knowledge and real-world application. Furthermore, the "
    "university's state-of-the-art facilities and diverse, supportive academic "
    "community provide the ideal environment for my professional growth. I am "
    "confident that the skills and insights gained from this course will empower "
    "me to make a meaningful contribution to the field."
)
_ESSAY_2412 = (
    "I chose the University of Hertfordshire because of its outstanding reputation "
    "as an innovative, business-facing university that prioritizes student "
    "employability. The university's strong corporate partnerships and historical "
    "links with industry leaders offer invaluable networking and placement "
    "opportunities. I am particularly drawn to UH's modern campus facilities, "
    "exceptional student support services, and its vibrant, multicultural academic "
    "environment. Additionally, its strategic location near London provides the "
    "perfect balance of a focused, resource-rich study environment with easy access "
    "to one of the world's major economic and cultural hubs. I believe that UH's "
    "practical approach to education and dedication to developing future-ready "
    "professionals make it the ideal institution to support my academic and career "
    "aspirations."
)
_ESSAY_2413 = (
    "I am motivated to apply to a UK university due to the UK's world-renowned "
    "reputation for academic excellence and its innovative approach to higher "
    "education. The UK higher education system is distinguished by its rigorous "
    "standards and highly efficient, intensive course structures, which allow me "
    "to obtain a globally recognized degree while accelerating my career entry. "
    "Studying in the UK offers an unparalleled opportunity to immerse myself in a "
    "native English-speaking environment and engage with a diverse, international "
    "student community. This global exposure will not only broaden my perspective "
    "but also foster essential cross-cultural communication skills. I believe that "
    "the combination of high-quality research, practical teaching methods, and "
    "vibrant cultural heritage in the UK makes it the ideal destination to achieve "
    "my academic and personal potential."
)
_ESSAY_2414 = (
    "My immediate career aspiration is to secure a managerial or specialist role "
    "within a progressive, global organization, where I can drive strategic growth "
    "and project efficiency. In the long term, I aim to transition into a senior "
    "leadership position, shaping sustainable business strategies in an increasingly "
    "interconnected international market. To achieve this, I need to complement my "
    "current background with advanced industry knowledge and refined analytical "
    "skills. This course perfectly bridges that gap. Its comprehensive curriculum "
    "covers critical modules that directly address complex market challenges. "
    "Furthermore, the university's emphasis on practical learning and "
    "industry-facing projects will allow me to apply theoretical frameworks to "
    "real-world scenarios. I am confident that the insights, global perspective, "
    "and practical expertise gained from this program will serve as the ideal "
    "catalyst for my professional career."
)

# 与 G 列同一行：B=2, AF=32, AG=33, AI=35, AJ=36（iter_rows min_col=2 时的下标）
_IDX_B = 2 - 2
_IDX_AJ = 36 - 2
# 与 G 列同一行：AF=32, AG=33, AI=35（iter_rows min_col=7 时的下标）
_IDX_AF = 32 - 7
_IDX_AG = 33 - 7
_IDX_AI = 35 - 7

_DEFAULT_XLSX = Path.home() / "Desktop" / "学生基础信息_已翻译.xlsx"

_COOKIE_ACCEPT_NAMES = (
    "Accept all",
    "Accept All",
    "Accept",
    "I agree",
    "同意",
    "接受全部",
)

# Cloudflare 人机验证常见文案（无法程序化绕过，只能手动点完后继续）
_CF_TEXT_HINTS = (
    "Verify you are human",
    "Just a moment",
    "Checking your browser",
    "确认您是真人",
    "请完成以下操作",
)
_CF_SELECTORS = (
    "#challenge-running",
    "#cf-challenge-running",
    "iframe[src*='challenges.cloudflare']",
    ".cf-turnstile",
)

_DEFAULT_USER_DATA = Path.home() / ".herts-playwright-profile"


def _excel_cell_str(raw: object) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    return "" if s.lower() == "nan" else s


def _gender_label_from_cell(raw: object) -> str:
    """J 列性别 → 网申选项 Male / Female（与 UEL 脚本列含义一致）。"""
    if raw is None:
        raise ValueError("性别为空")
    s = str(raw).strip()
    if not s or s.lower() == "nan":
        raise ValueError("性别为空")
    key = s.lower().replace(" ", "")
    if key in {"f", "female", "woman", "女", "女性"} or key.startswith("female"):
        return "Female"
    if key in {"m", "male", "man", "男", "男性"} or key.startswith("male"):
        return "Male"
    raise ValueError(f"无法识别性别 {raw!r}，请填写 Male/Female、M/F 或 男/女")


def student_identity_from_xlsx(
    xlsx_path: Path,
) -> tuple[str, str, str, str, str, str]:
    """G 列姓名 + I 列身份证 + J 列性别 → 姓、名、月、日、年、Male/Female。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=10, values_only=True):
            g_cell, _h, i_cell, j_cell = row[0], row[1], row[2], row[3]
            if g_cell is None:
                continue
            full = str(g_cell).strip()
            if not full:
                continue
            parts = full.split(None, 1)
            surname, forename = (
                (parts[0].strip(), parts[1].strip()) if len(parts) > 1 else ("", full)
            )
            ymd = _dob_yyyymmdd_from_chinese_id(i_cell)
            dt = datetime.strptime(ymd, "%Y%m%d")
            gender = _gender_label_from_cell(j_cell)
            return (
                surname,
                forename,
                str(dt.month),
                str(dt.day),
                str(dt.year),
                gender,
            )
        raise ValueError(f"{xlsx_path} 的 G/I/J 列（自第 2 行起）没有可用学生信息")
    finally:
        wb.close()


def _phone_from_cell(raw: object) -> str:
    if raw is None:
        return ""
    if isinstance(raw, float):
        if math.isnan(raw):
            return ""
        if raw.is_integer():
            return str(int(raw))
        return str(raw).strip()
    if isinstance(raw, int) and not isinstance(raw, bool):
        return str(raw)
    return _excel_cell_str(raw)


def _normalize_cn_mobile(phone: str) -> str:
    """AI 列手机号 → 11 位国内号码（intl-tel-input 在 +86 下填写）。"""
    digits = re.sub(r"\D", "", phone)
    if digits.startswith("86") and len(digits) >= 13:
        digits = digits[2:]
    if digits.startswith("0086"):
        digits = digits[4:]
    if len(digits) == 12 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) != 11:
        raise ValueError(f"手机号须为 11 位（当前 {len(digits)} 位）: {phone!r}")
    return digits


def _select_phone_country_china(scope: FrameLocator) -> None:
    """intl-tel-input：选择 China (中国) +86。"""
    flag = scope.locator(_PHONE_FLAG)
    if flag.count() == 0:
        flag = scope.locator(_PHONE_FLAG_FALLBACK)
    flag.first.wait_for(state="visible", timeout=60_000)

    title = (flag.first.get_attribute("title") or "").strip()
    if "+86" in title or "China" in title or "中国" in title:
        print("-> 手机区号已是 China +86")
        return

    flag.first.click()
    print("-> 已展开手机区号列表")

    for sel in (
        f"{_PHONE_LISTBOX} li",
        '[id="form_1551_listbox"] li',
        ".iti__country-list li",
    ):
        options = scope.locator(sel)
        if options.count() == 0:
            continue
        china_loc = options.filter(has_text=re.compile(r"China|\+86|中国", re.I))
        if china_loc.count() > 0:
            china_loc.first.scroll_into_view_if_needed()
            china_loc.first.click()
            print("-> 已选择手机区号: China +86")
            return

    china_opt = scope.get_by_role("option", name=re.compile(r"China|\+86|中国", re.I))
    if china_opt.count() > 0:
        china_opt.first.click()
        print("-> 已选择手机区号: China +86（role=option）")
        return

    raise RuntimeError("未能在区号列表中找到 China (+86)")


def _fill_international_phone(scope: FrameLocator, phone: str) -> None:
    """填写 #form_1551：先选 +86，再填 AI 列手机号。"""
    digits = _normalize_cn_mobile(phone)
    _select_phone_country_china(scope)

    tel = scope.locator(_PHONE)
    tel.wait_for(state="visible", timeout=60_000)
    tel.scroll_into_view_if_needed()
    tel.click()
    tel.fill("")
    tel.fill(digits)
    tel.press("Tab")

    # 等待校验通过（去掉 ng-invalid-international-phone-number）
    try:
        scope.locator(f"{_PHONE}.ng-invalid-international-phone-number").wait_for(
            state="hidden", timeout=8_000
        )
    except Exception:
        cls = tel.get_attribute("class") or ""
        if "ng-invalid-international-phone-number" in cls:
            print(
                f"[警告] 手机号仍校验失败 class={cls!r}，"
                f"已填号码={digits}，请检查 AI 列或页面格式要求"
            )
        else:
            print(f"-> 已填写手机（+86）: {digits}")
            return
    print(f"-> 已填写手机（+86）: {digits}")


def _parse_af_for_herts(af_raw: str) -> tuple[str, str, str]:
    """AF 常住址 → 省、市、街道。例：浙江省,宁波市,奉化区;详细地址"""
    region, street = af_raw, ""
    if ";" in af_raw:
        region, street = af_raw.split(";", 1)
        region, street = region.strip(), street.strip()
    parts = [p.strip() for p in region.split(",") if p.strip()]
    province = parts[0] if parts else ""
    city = parts[1] if len(parts) > 1 else ""
    if not street:
        street = ", ".join(parts[2:]) if len(parts) > 2 else "" 
    elif len(parts) > 2 and parts[2]:
        district = parts[2]
        if district not in street:
            street = f"{district}, {street}" if street else district
    if not province or not city or not street:
        raise ValueError(f"AF 列地址无法拆成省/市/街道: {af_raw!r}")
    return province, city, street


def contact_from_xlsx(xlsx_path: Path) -> tuple[str, str, str, str, str]:
    """与 G 列同一行：AI 手机；AF 省/市/街道；AG 邮编。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=35, values_only=True):
            g_cell = row[0]
            if g_cell is None:
                continue
            if not str(g_cell).strip():
                continue
            af_raw = _excel_cell_str(row[_IDX_AF])
            postcode = _excel_cell_str(row[_IDX_AG])
            phone = _phone_from_cell(row[_IDX_AI])
            if not af_raw:
                raise ValueError(f"{xlsx_path} 的 AF 列（地址）为空")
            if not postcode:
                raise ValueError(f"{xlsx_path} 的 AG 列（邮编）为空")
            if not phone:
                raise ValueError(f"{xlsx_path} 的 AI 列（手机）为空")
            province, city, street = _parse_af_for_herts(af_raw)
            return phone, street, city, province, postcode
        raise ValueError(f"{xlsx_path} 的 G/AF/AG/AI 列（自第 2 行起）没有可用联系信息")
    finally:
        wb.close()


def major_and_aj_from_xlsx(xlsx_path: Path) -> tuple[str, str]:
    """与 G 列同一行：B 列（专业）、AJ 列。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        g_offset = 7 - 2  # G 列在 min_col=2 时的下标
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=36, values_only=True):
            g_cell = row[g_offset] if len(row) > g_offset else None
            if g_cell is None:
                continue
            if not str(g_cell).strip():
                continue
            major = _excel_cell_str(row[_IDX_B])
            aj_val = _excel_cell_str(row[_IDX_AJ]) if len(row) > _IDX_AJ else ""
            if not major:
                raise ValueError(f"{xlsx_path} 的 B 列（MAJOR）为空")
            if not aj_val:
                raise ValueError(f"{xlsx_path} 的 AJ 列为空")
            return major, aj_val
        raise ValueError(f"{xlsx_path} 的 B/AJ 列（自第 2 行起）没有可用数据")
    finally:
        wb.close()


def _table_cell_input(col: int) -> str:
    return (
        f"#field2474 repeatable-table-field tbody tr td:nth-child({col}) input"
    )


def _table_cell_select(col: int) -> str:
    return (
        f"#field2474 repeatable-table-field tbody tr td:nth-child({col}) select"
    )


def _resolve_desktop_upload_path(
    *,
    env_keys: tuple[str, ...],
    default_path: Path,
    glob_stems: tuple[str, ...],
    doc_label: str,
) -> Path:
    """解析桌面待上传文件：环境变量 > 默认路径 > 桌面 glob。"""
    for key in env_keys:
        env = os.environ.get(key, "").strip()
        if env:
            path = Path(env).expanduser()
            if path.is_file():
                return path.resolve()
            raise FileNotFoundError(f"{key} 指向的文件不存在: {path}")

    if default_path.is_file():
        return default_path.resolve()

    desktop = Path.home() / "Desktop"
    for stem in glob_stems:
        for ext in (".pdf", ".PDF", ".png", ".jpg", ".jpeg"):
            cand = desktop / f"{stem}{ext}"
            if cand.is_file():
                return cand.resolve()
        matches = sorted(
            (p for p in desktop.glob(f"{stem}*") if p.is_file()),
            key=lambda p: p.name.lower(),
        )
        if matches:
            return matches[0].resolve()

    raise FileNotFoundError(
        f"找不到{doc_label}。请将文件放在桌面（默认: {default_path}），"
        f"或设置环境变量 {env_keys[0]}"
    )


def _resolve_transcript_path() -> Path:
    """成绩单：环境变量 > 桌面 transcript.pdf > transcript*。"""
    return _resolve_desktop_upload_path(
        env_keys=("HERTS_TRANSCRIPT_PDF", "HERTS_TRANSCRIPT", "UEL_TRANSCRIPT_PDF"),
        default_path=_DEFAULT_TRANSCRIPT,
        glob_stems=("transcript",),
        doc_label="成绩单 transcript",
    )


def _resolve_certificate_path() -> Path:
    """证书：环境变量 > 桌面 certficate.pdf（与 UEL 一致）> certificate*。"""
    return _resolve_desktop_upload_path(
        env_keys=(
            "HERTS_CERTIFICATE_PDF",
            "HERTS_CERTIFICATE",
            "UEL_CERTIFICATE_PDF",
        ),
        default_path=_DEFAULT_CERTIFICATE,
        glob_stems=("certficate", "certificate"),
        doc_label="证书 certificate",
    )


def _scroll_upload_section_into_view(frame: FrameLocator) -> None:
    """第 5 页上传区滚入视口一次即可，避免每次点击都 scroll 导致页面上下跳。"""
    section = frame.locator("#field3144")
    section.wait_for(state="visible", timeout=60_000)
    section.scroll_into_view_if_needed()


def _choose_file_in_herts_modal(
    frame: FrameLocator,
    page: Page,
    file_path: Path,
    label: str,
) -> None:
    """直接通过隐藏 input 选文件，等待文件名出现在弹窗待传列表。"""
    path = file_path.expanduser()
    resolved = str(path.resolve())

    hidden_input = frame.locator(_FIELD3144_INPUT_FILE)
    hidden_input.set_input_files(resolved)
    print(f"-> {label} 已通过 input 域快速选择: {path.name}")

    modal = frame.locator("#field3144 .modal")
    modal.get_by_text(path.name, exact=False).first.wait_for(
        state="visible", timeout=15_000
    )
    print(f"-> {label} 已出现在弹窗待传列表")


def _wait_herts_modal_upload_done(
    frame: FrameLocator, page: Page, *, doc_label: str, timeout_ms: int = 120_000
) -> None:
    """点击弹窗 Upload 后等待服务器传完（弹窗关闭）。期间页面可能不动，属正常上传。"""
    modal = frame.locator("#field3144 .modal")
    confirm_btn = frame.locator(_FILE_UPLOAD_CONFIRM)

    for _ in range(20):
        try:
            aria = confirm_btn.get_attribute("aria-disabled") or ""
            ng = confirm_btn.get_attribute("ng-disabled") or ""
            if aria.lower() == "true" or ng.lower() == "true":
                break
        except Exception:
            pass
        page.wait_for_timeout(300)

    try:
        modal.wait_for(state="hidden", timeout=timeout_ms)
        print(f"-> {doc_label} 上传完成（弹窗已关闭）")
        return
    except Exception:
        print(f"[提示] {doc_label} 等待弹窗关闭超时，继续下一步")


def _wait_herts_uploads_ready_for_next(
    frame: FrameLocator, page: Page, *, timeout_ms: int = 120_000
) -> None:
    """两个文件都传完后，等待 #form-next 可点（上传不再阻塞）。"""
    next_btn = frame.locator(_FORM_NEXT)
    next_btn.wait_for(state="visible", timeout=60_000)
    for _ in range(timeout_ms // 500):
        try:
            if next_btn.is_enabled():
                print("-> 上传已全部完成，Next 可点击")
                return
        except Exception:
            pass
        page.wait_for_timeout(500)
    print("[警告] Next 仍为 disabled，将尝试强制点击翻页")


def _upload_herts_field3144_file(
    frame: FrameLocator, page: Page, file_path: Path, label: str
) -> None:
    """第 5 页：逐个上传。开弹窗 → Choose a file → 列表出现 → 点 Upload → 等完成。"""
    path = file_path.expanduser()
    if not path.is_file():
        raise FileNotFoundError(f"找不到{label}: {path}")

    toggle_btn = frame.locator(_FIELD3144_TOGGLE_UPLOAD)
    toggle_btn.wait_for(state="visible", timeout=60_000)
    toggle_btn.click()
    print(f"-> 已点击 Upload 打开弹窗（{label}）")

    frame.locator("#field3144 .modal").wait_for(state="visible", timeout=30_000)
    _choose_file_in_herts_modal(frame, page, path, label)

    confirm_btn = frame.locator(_FILE_UPLOAD_CONFIRM)
    confirm_btn.wait_for(state="visible", timeout=30_000)
    try:
        expect(confirm_btn).to_be_enabled(timeout=15_000)
    except Exception:
        page.wait_for_timeout(800)
    confirm_btn.click()
    print("-> 已点击弹窗内 Upload（#file-upload_3144），等待服务器上传…")

    _wait_herts_modal_upload_done(frame, page, doc_label=label)
    page.wait_for_timeout(300)


def _fill_select_option(
    scope: FrameLocator, selector: str, option_text: str, label: str
) -> None:
    sel = scope.locator(selector)
    sel.wait_for(state="visible", timeout=60_000)
    sel.scroll_into_view_if_needed()
    try:
        sel.select_option(label=option_text)
    except Exception:
        try:
            sel.select_option(value=f"string:{option_text}")
        except Exception:
            sel.select_option(value=option_text)
    print(f"-> 已选择 {label}: {option_text}")


def _cloudflare_challenge_visible(page: Page) -> bool:
    url = page.url.lower()
    if "challenges.cloudflare" in url:
        return True
    for sel in _CF_SELECTORS:
        loc = page.locator(sel)
        if loc.count() == 0:
            continue
        try:
            if loc.first.is_visible():
                return True
        except Exception:
            continue
    for hint in _CF_TEXT_HINTS:
        loc = page.get_by_text(hint)
        if loc.count() == 0:
            continue
        try:
            if loc.first.is_visible():
                return True
        except Exception:
            continue
    return False


def _wait_cloudflare_manual(page: Page, step: str = "") -> None:
    """若出现 Cloudflare，暂停脚本，由用户在浏览器里手动通过验证。"""
    if not _cloudflare_challenge_visible(page):
        return
    where = f"（{step}）" if step else ""
    print(
        f"\n[Cloudflare] 检测到人机验证{where}。"
        "无法自动跳过，请在浏览器窗口中手动完成验证。"
    )
    input(">>> 验证通过后，在此按 Enter 继续运行脚本…")
    page.wait_for_load_state("domcontentloaded", timeout=120_000)
    try:
        page.wait_for_load_state("networkidle", timeout=30_000)
    except Exception:
        pass


def _dismiss_ai_robot_overlay(page: Page, frame: FrameLocator) -> None:
    """第 1 页填完后、点 Next 前：点击 fa-xmark 关闭 AI 机器人浮层。"""
    scopes: list[tuple[str, Page | FrameLocator]] = [("主页面", page), ("iframe", frame)]
    selectors = (
        ('button:has(svg[data-icon="xmark"])', "xmark 按钮"),
        ('button:has(.fa-xmark)', "fa-xmark 按钮"),
        ('a:has(.fa-xmark)', "fa-xmark 链接"),
        ("svg.fa-xmark", "xmark 图标"),
        ('svg[data-icon="xmark"]', "xmark svg"),
    )

    for scope_label, root in scopes:
        for sel, desc in selectors:
            loc = root.locator(sel)
            for i in range(loc.count()):
                target = loc.nth(i)
                try:
                    if not target.is_visible():
                        continue
                    if sel.startswith("svg"):
                        parent = target.locator("xpath=ancestor::button[1]")
                        if parent.count() > 0:
                            parent.first.click()
                        else:
                            target.click(force=True)
                    else:
                        target.click()
                    print(f"-> 已关闭 AI 浮层（{scope_label} {desc}）")
                    page.wait_for_timeout(600)
                    return
                except Exception:
                    continue

    print("[提示] 未自动找到 xmark 关闭钮，若 AI 浮层仍在请手动关闭后再点 Next")


def _dismiss_cookie_banner_if_present(page: Page) -> None:
    for name in _COOKIE_ACCEPT_NAMES:
        btn = page.get_by_role("button", name=name)
        if btn.count() == 0:
            continue
        for i in range(btn.count()):
            candidate = btn.nth(i)
            if candidate.is_visible():
                candidate.click()
                print(f"-> 已关闭 Cookie 提示（{name!r}）")
                page.wait_for_timeout(500)
                return


def _form_frame(page: Page) -> FrameLocator:
    """定位申请表单 iframe（排除 OneTrust 的 ot-text-resize iframe）。"""
    sel = _FORM_IFRAME if page.locator(_FORM_IFRAME).count() > 0 else _FORM_IFRAME_FALLBACK
    page.locator(sel).wait_for(state="attached", timeout=120_000)
    return page.frame_locator(sel)


def _get_form_page_progress(scope: FrameLocator) -> str:
    """读取页眉进度，例如 Page 1 / 6。"""
    loc = scope.locator("div.page-count.text-center")
    if loc.count() == 0:
        loc = scope.locator("div.page-count")
    if loc.count() == 0:
        return ""
    return re.sub(r"\s+", " ", loc.first.inner_text().strip())


def _log_form_page_progress(scope: FrameLocator, step: str) -> None:
    text = _get_form_page_progress(scope)
    if text:
        print(f"[进度] {step}: {text}")
    else:
        print(f"[进度] {step}: （未检测到 page-count，可能尚未进入表单或 DOM 已变）")


def _find_next_button(scope: FrameLocator, page: Page):
    """在 iframe 或主页面中定位 Next 按钮（#form-next / XPath）。"""
    candidates: list[tuple[str, object]] = [
        ("iframe #form-next", scope.locator(_FORM_NEXT)),
        ("iframe xpath", scope.locator(_FORM_NEXT_XPATH)),
        ("iframe aria-label", scope.locator('button[aria-label="Next"]')),
        ("iframe 按钮 Next", scope.get_by_role("button", name=re.compile(r"^Next$", re.I))),
        ("主页面 #form-next", page.locator(_FORM_NEXT)),
        ("主页面 xpath", page.locator(_FORM_NEXT_XPATH)),
        ("主页面 aria-label", page.locator('button[aria-label="Next"]')),
        ("主页面 按钮 Next", page.get_by_role("button", name=re.compile(r"^Next$", re.I))),
    ]
    for source, loc in candidates:
        if loc.count() == 0:
            continue
        btn = loc.first
        try:
            if btn.is_visible():
                return btn, source
        except Exception:
            continue
    return None, ""


def _trigger_next_via_angular(scope: FrameLocator) -> str | None:
    """直接调用 ng-click 绑定的 changePage(currentPage + 1, true)。"""
    anchor = scope.locator("#form-next")
    if anchor.count() == 0:
        return None
    try:
        result = anchor.first.evaluate(
            """(btn) => {
                if (!btn) return { ok: false, reason: 'no #form-next' };
                try {
                    if (typeof angular !== 'undefined') {
                        let s = angular.element(btn).scope();
                        for (let i = 0; i < 10 && s && typeof s.changePage !== 'function'; i++) {
                            s = s.$parent;
                        }
                        if (s && typeof s.changePage === 'function') {
                            const next = (s.currentPage || 0) + 1;
                            s.$apply(() => s.changePage(next, true));
                            return { ok: true, method: 'changePage', next };
                        }
                    }
                } catch (e) {
                    return { ok: false, reason: 'angular: ' + e.message };
                }
                btn.disabled = false;
                btn.removeAttribute('disabled');
                btn.removeAttribute('ng-disabled');
                btn.click();
                return { ok: true, method: 'dom-click' };
            }"""
        )
    except Exception:
        return None
    if isinstance(result, dict) and result.get("ok"):
        return str(result.get("method", "angular"))
    if isinstance(result, dict):
        print(f"[提示] Angular 翻页失败: {result.get('reason')}")
    return None


def _page_reached(scope: FrameLocator, expect_page: int) -> bool:
    text = _get_form_page_progress(scope)
    return bool(re.search(rf"Page\s+{expect_page}\s*/", text, re.I))


def _click_form_next(
    scope: FrameLocator, page: Page, expect_page: int | None = None
) -> None:
    """点击 Next（#form-next），含 XPath 与 Angular changePage。"""
    before = _get_form_page_progress(scope)

    btn, source = _find_next_button(scope, page)
    if btn is None:
        raise RuntimeError("未找到可见的 Next 按钮（#form-next / //*[@id='form-next']）")

    btn.wait_for(state="visible", timeout=60_000)
    btn.scroll_into_view_if_needed()

    for _ in range(40):
        try:
            if btn.is_enabled():
                break
        except Exception:
            break
        page.wait_for_timeout(500)
    else:
        print(
            "[警告] Next 仍为 disabled（可能 ng-disabled=isSaveButtonDisabledDueToUpload "
            "或表单校验未通过），将强制尝试翻页…"
        )

    print(f"-> 准备点击 Next（{source}）")

    if expect_page is not None and _page_reached(scope, expect_page):
        _log_form_page_progress(scope, "已在目标页")
        return

    def _do_angular() -> None:
        if not _trigger_next_via_angular(scope):
            raise RuntimeError("angular changePage 未执行")

    for name, action in (
        ("click", lambda: btn.click(timeout=15_000)),
        ("force-click", lambda: btn.click(force=True, timeout=15_000)),
        ("angular-changePage", _do_angular),
    ):
        try:
            action()
            print(f"-> 已执行 Next（{name}）")
            page.wait_for_timeout(1200)
            if expect_page is None or _page_reached(scope, expect_page):
                if expect_page is not None:
                    _log_form_page_progress(scope, "翻页后")
                return
        except Exception as e:
            print(f"[提示] Next 方式 {name!r} 未翻页: {e}")

    if expect_page is not None:
        pat = re.compile(rf"Page\s+{expect_page}\s*/\s*\d+", re.I)
        try:
            scope.locator("div.page-count").filter(has_text=pat).wait_for(
                state="visible", timeout=15_000
            )
            _log_form_page_progress(scope, "翻页后")
            return
        except Exception:
            after = _get_form_page_progress(scope)
            print(f"[提示] 翻页前: {before!r}，翻页后: {after!r}")
            print(
                f"[提示] 未进入 Page {expect_page}。请检查：\n"
                "  1) 第 1 页红色校验错误\n"
                "  2) Next 是否灰色（isSaveButtonDisabledDueToUpload 上传未完成）\n"
                "  3) 在浏览器中手动点一次 Next 确认能否翻页"
            )
            raise RuntimeError(f"未能翻到第 {expect_page} 页")


def _fill_text(
    scope: FrameLocator,
    selector: str,
    text: str,
    label: str,
    *,
    log_preview: int | None = None,
) -> None:
    field = scope.locator(selector)
    field.wait_for(state="visible", timeout=60_000)
    field.fill(text)
    if log_preview is not None and len(text) > log_preview:
        shown = f"{text[:log_preview]}…（共 {len(text)} 字符）"
    else:
        shown = text
    print(f"-> 已填写 {label}: {shown}")


def _combo_container_id(combobox_id: str) -> str:
    return combobox_id.replace("-combobox-container", "_container")


def _combobox_id_from_container(container_id: str) -> str:
    base = container_id.removesuffix("_container")
    return f"{base}-combobox-container"


def _select_chosen_by_container(
    scope: FrameLocator,
    container_id: str,
    option_text: str,
    label: str,
    *,
    option_alternates: tuple[str, ...] = (),
) -> None:
    """通过 #form_XXXX_container > a 选择 Chosen 项。"""
    _select_chosen_option(
        scope,
        combobox_id=_combobox_id_from_container(container_id),
        container_id=container_id,
        option_text=option_text,
        label=label,
        option_alternates=option_alternates,
    )


def _select_chosen_option(
    scope: FrameLocator,
    *,
    combobox_id: str,
    option_text: str,
    label: str,
    container_id: str | None = None,
    option_alternates: tuple[str, ...] = (),
) -> None:
    """Chosen 下拉：点「Select an Option」展开，再选指定项。"""
    container_id = container_id or _combo_container_id(combobox_id)
    texts = (option_text, *option_alternates)

    try:
        scope.locator("body").press("Escape")
    except Exception:
        pass

    toggle = None
    for sel in (
        f'a.chosen-single[aria-controls="{combobox_id}"]',
        f"#{container_id} > a.chosen-single",
        f"#{container_id} > a",
    ):
        loc = scope.locator(sel)
        if loc.count() > 0:
            toggle = loc.first
            break
    if toggle is None:
        raise RuntimeError(f"未找到 {label} 下拉按钮: {combobox_id}")

    toggle.wait_for(state="visible", timeout=60_000)
    toggle.scroll_into_view_if_needed()
    toggle.click()
    print(f"-> 已展开 {label} 下拉（{container_id}）")

    expanded = scope.locator(
        f'a.chosen-single[aria-controls="{combobox_id}"][aria-expanded="true"]'
    )
    try:
        expanded.wait_for(state="attached", timeout=5_000)
    except Exception:
        toggle.click(force=True)
        expanded.wait_for(state="attached", timeout=10_000)

    for search_sel in (
        f"#{container_id} .chosen-search input",
        f"#{combobox_id} .chosen-search input",
    ):
        search = scope.locator(search_sel)
        if search.count() > 0 and search.first.is_visible():
            search.first.fill(option_text)
            break

    exact = re.compile(rf"^{re.escape(option_text)}$")
    option = None
    result_scopes = (
        f"#{combobox_id} .chosen-results li",
        f"#{container_id} .chosen-results li",
        ".chosen-container-active .chosen-results li",
        ".chosen-drop:visible .chosen-results li",
    )
    for text in texts:
        pat = re.compile(rf"^{re.escape(text)}$")
        for rs in result_scopes:
            loc = scope.locator(rs).filter(has_text=pat)
            if loc.count() > 0:
                option = loc.first
                break
            loc = scope.locator(rs).filter(has=scope.locator("span", has_text=text))
            if loc.count() > 0:
                option = loc.first
                break
        if option is not None:
            break

    if option is None:
        option = scope.locator(".chosen-drop:visible .chosen-results li").filter(
            has_text=exact
        ).first

    option.wait_for(state="visible", timeout=30_000)
    option.scroll_into_view_if_needed()
    try:
        option.click(timeout=10_000)
    except Exception:
        option.click(force=True)
    print(f"-> 已选择 {label}: {option_text}")


def _fill_personal_details_form(page: Page, xlsx: Path) -> None:
    surname, forename, month, day, year, gender = student_identity_from_xlsx(xlsx)
    phone, street, city, province, postcode = contact_from_xlsx(xlsx)
    print(
        f"[数据] 姓={surname!r} 名={forename!r} 性别={gender} "
        f"生日={month}/{day}/{year} 手机={phone}"
    )
    print(f"[数据] 地址: {province} / {city} / {street} / 邮编 {postcode}")

    frame = _form_frame(page)
    frame.locator(_FIRST_NAME).wait_for(state="visible", timeout=120_000)
    _log_form_page_progress(frame, "开始填写")

    _select_chosen_option(
        frame, combobox_id=_TITLE_COMBO, option_text="Mx", label="称谓"
    )
    _select_chosen_option(
        frame, combobox_id=_GENDER_COMBO, option_text=gender, label="性别"
    )
    _fill_text(frame, _FIRST_NAME, forename, "名 First name")
    _fill_text(frame, _LAST_NAME, surname, "姓 Last name")
    _fill_text(frame, _EMAIL, email, "Email")
    _fill_text(frame, _EMAIL_CONFIRM, email, "确认 Email")

    _select_chosen_option(
        frame, combobox_id=_AGE_18_COMBO, option_text="No", label="满18岁确认"
    )

    _fill_text(frame, _DOB_MONTH, month, "出生月")
    _fill_text(frame, _DOB_DAY, day, "出生日")
    _fill_text(frame, _DOB_YEAR, year, "出生年")

    _select_chosen_option(
        frame, combobox_id=_COUNTRY_COMBO, option_text="China", label="居住地"
    )
    _select_chosen_option(
        frame, combobox_id=_NATIONALITY_COMBO, option_text="Chinese", label="国籍"
    )
    _select_chosen_option(
        frame,
        combobox_id=_BIRTH_COUNTRY_COMBO,
        container_id="form_17173_container",
        option_text="China",
        option_alternates=("CHINA", "China, People's Republic of"),
        label="出生地",
    )
    _select_chosen_option(
        frame, combobox_id=_CRIMINAL_COMBO, option_text="No", label="犯罪记录"
    )
    _select_chosen_option(
        frame, combobox_id=_DISABILITY_COMBO, option_text="No", label="残疾情况"
    )
    _select_chosen_option(
        frame, combobox_id=_FORM_2362_COMBO, option_text="No", label="form_2362"
    )

    _fill_international_phone(frame, phone)
    _fill_text(frame, _STREET, street, "街道地址")
    _fill_text(frame, _CITY, city, "城市")
    _fill_text(frame, _PROVINCE, province, "省")
    _fill_text(frame, _POSTCODE, postcode, "邮编")
    frame.locator(_POSTCODE).press("Tab")
    page.wait_for_timeout(800)

    _log_form_page_progress(frame, "第 1 页填写结束")
    print("[完成] 第 1 / 6 页个人基本信息已填写")

    _dismiss_ai_robot_overlay(page, frame)
    _click_form_next(frame, page, expect_page=2)
    _fill_page_2(frame, page)
    _fill_page_3(frame, page, xlsx)
    _skip_page_4(frame, page)
    _fill_page_5(frame, page)
    _fill_page_6(frame, page)


def _fill_page_2(frame: FrameLocator, page: Page) -> None:
    """第 2 页：入学时间、课程及若干 No 选项。"""
    _log_form_page_progress(frame, "开始填写第 2 页")
    if not _page_reached(frame, 2):
        raise RuntimeError("未进入 Page 2，无法填写第 2 页")

    page.wait_for_timeout(500)

    _select_chosen_by_container(
        frame, "form_1573_container", "September 2026", "入学时间"
    )
    _select_chosen_by_container(
        frame,
        "form_2819_container",
        "MA Human Resource Management",
        "课程",
    )

    for container_id, label in (
        ("form_2384_container", "form_2384"),
        ("form_2385_container", "form_2385"),
        ("form_2387_container", "form_2387"),
        ("form_2389_container", "form_2389"),
        ("form_2390_container", "form_2390"),
    ):
        _select_chosen_by_container(
            frame,
            container_id,
            "No",
            label,
            option_alternates=("NO",),
        )

    _log_form_page_progress(frame, "第 2 页填写结束")
    print("[完成] 第 2 / 6 页已填写")

    _click_form_next(frame, page, expect_page=3)


def _fill_page_3(frame: FrameLocator, page: Page, xlsx: Path) -> None:
    """第 3 页：学校、教育经历表格、form_2399、个人陈述四题。"""
    _log_form_page_progress(frame, "开始填写第 3 页")
    if not _page_reached(frame, 3):
        raise RuntimeError("未进入 Page 3，无法填写第 3 页")

    major, aj_val = major_and_aj_from_xlsx(xlsx)
    print(f"[数据] B 列={major!r} AJ 列={aj_val!r}")

    page.wait_for_timeout(500)
    frame.locator(_TABLE_ROW).first.wait_for(state="visible", timeout=60_000)

    _fill_text(frame, _FORM_3683, _INSTITUTION, "学校 form_3683")
    _fill_text(frame, _table_cell_input(1), _INSTITUTION, "表-学校")
    _fill_text(frame, _table_cell_input(2), "China", "表-国家")
    _fill_text(frame, _table_cell_input(3), "China", "表-国家2")
    _fill_text(frame, _table_cell_input(4), major, "表-B列")
    _fill_text(frame, _table_cell_input(5), aj_val, "表-AJ列")
    _fill_select_option(
        frame, _table_cell_select(6), "Pending", "表-状态",
    )

    _select_chosen_by_container(
        frame, _FORM_2399_CONTAINER, "None", "form_2399"
    )
    _fill_text(frame, _FORM_2411, _ESSAY_2411, "form_2411", log_preview=80)
    _fill_text(frame, _FORM_2412, _ESSAY_2412, "form_2412", log_preview=80)
    _fill_text(frame, _FORM_2413, _ESSAY_2413, "form_2413", log_preview=80)
    _fill_text(frame, _FORM_2414, _ESSAY_2414, "form_2414", log_preview=80)

    _select_chosen_by_container(
        frame,
        _FORM_3173_CONTAINER,
        "No",
        "form_3173",
        option_alternates=("NO",),
    )

    _log_form_page_progress(frame, "第 3 页填写结束")
    _click_form_next(frame, page, expect_page=4)
    print("[完成] 第 3 / 6 页已填写")


def _skip_page_4(frame: FrameLocator, page: Page) -> None:
    """第 4 页无字段，直接点 Next 进入第 5 页。"""
    _log_form_page_progress(frame, "第 4 页（无字段，直接 Next）")
    if not _page_reached(frame, 4):
        raise RuntimeError("未进入 Page 4，无法跳过第 4 页")
    page.wait_for_timeout(500)
    _click_form_next(frame, page, expect_page=5)
    print("[完成] 已通过第 4 / 6 页（空页）")


def _fill_page_5(frame: FrameLocator, page: Page) -> None:
    """第 5 页：是否上传选 Yes，依次上传 transcript 与 certificate。"""
    _log_form_page_progress(frame, "开始填写第 5 页")
    if not _page_reached(frame, 5):
        raise RuntimeError("未进入 Page 5，无法填写第 5 页")

    page.wait_for_timeout(500)

    _select_chosen_by_container(
        frame,
        _FORM_2582_CONTAINER,
        "Yes",
        "form_2582",
        option_alternates=("YES",),
    )
    page.wait_for_timeout(500)
    frame.locator(_FIELD3144_TOGGLE_UPLOAD).wait_for(
        state="visible", timeout=60_000
    )
    _scroll_upload_section_into_view(frame)

    transcript = _resolve_transcript_path()
    certificate = _resolve_certificate_path()
    print(f"[数据] 成绩单: {transcript}")
    print(f"[数据] 证书: {certificate}")

    _upload_herts_field3144_file(
        frame, page, transcript, "成绩单 transcript"
    )
    _upload_herts_field3144_file(
        frame, page, certificate, "证书 certificate"
    )

    _wait_herts_uploads_ready_for_next(frame, page)
    _click_form_next(frame, page, expect_page=6)

    _log_form_page_progress(frame, "第 5 页填写结束")
    print("[完成] 第 5 / 6 页已上传并进入第 6 页")


def _fill_page_6(frame: FrameLocator, page: Page) -> None:
    """第 6 页：条款同意 form_1602 → Yes, I agree to the above terms。"""
    _log_form_page_progress(frame, "开始填写第 6 页")
    if not _page_reached(frame, 6):
        raise RuntimeError("未进入 Page 6，无法填写第 6 页")

    page.wait_for_timeout(500)

    _select_chosen_by_container(
        frame,
        _FORM_1602_CONTAINER,
        _TERMS_AGREE_OPTION,
        "form_1602 条款同意",
    )

    _log_form_page_progress(frame, "第 6 页填写结束")
    print("[完成] 第 6 / 6 页已勾选条款同意")


def _safe_close_browser(context, browser) -> None:
    """关闭浏览器；若用户已手动关掉窗口，则忽略 TargetClosedError。"""
    try:
        context.close()
    except PlaywrightError as e:
        if "TargetClosed" in type(e).__name__ or "has been closed" in str(e):
            print("[提示] 浏览器已关闭，跳过 context.close()。")
        else:
            raise
    if browser is not None:
        try:
            browser.close()
        except PlaywrightError as e:
            if "TargetClosed" not in type(e).__name__ and "has been closed" not in str(e):
                raise


def _run_herts_apply_start(page: Page, xlsx: Path) -> None:
    print(f"[打开] {URL}")
    page.goto(URL, wait_until="domcontentloaded", timeout=120_000)
    page.wait_for_load_state("networkidle", timeout=60_000)
    _wait_cloudflare_manual(page, "打开申请页")
    _dismiss_cookie_banner_if_present(page)

    apply_btn = page.locator(_APPLY_BTN)
    apply_btn.wait_for(state="visible", timeout=60_000)
    apply_btn.scroll_into_view_if_needed()
    apply_btn.click()
    print("-> 已点击「申请」按钮（Make an application）")

    page.wait_for_load_state("domcontentloaded", timeout=60_000)
    page.wait_for_timeout(1500)
    _wait_cloudflare_manual(page, "进入 Gecko 申请表")

    _fill_personal_details_form(page, xlsx)


def main() -> None:
    xlsx = Path(os.environ.get("STUDENT_INFO_XLSX", str(_DEFAULT_XLSX))).expanduser()
    if not xlsx.is_file():
        raise FileNotFoundError(f"找不到学生信息表: {xlsx}")

    headless = os.getenv("HEADLESS", "").strip() in {"1", "true", "True", "yes", "YES"}
    use_chrome = os.getenv("HERTS_USE_CHROME", "1").strip() in {"1", "true", "True", "yes", "YES"}
    slow_mo = int(os.getenv("HERTS_SLOW_MO", "80") or "0")
    user_data = os.getenv("HERTS_USER_DATA_DIR", str(_DEFAULT_USER_DATA)).strip()

    launch_kw: dict = {"headless": headless}
    if use_chrome:
        launch_kw["channel"] = "chrome"
    if slow_mo > 0:
        launch_kw["slow_mo"] = slow_mo

    context_kw = {"locale": "en-GB", "timezone_id": "Europe/London"}

    with sync_playwright() as p:
        browser = None
        if user_data:
            profile = Path(user_data).expanduser()
            profile.mkdir(parents=True, exist_ok=True)
            print(f"[浏览器] 使用持久化配置（Cookie 可保留）: {profile}")
            context = p.chromium.launch_persistent_context(
                str(profile), **launch_kw, **context_kw
            )
            page = context.pages[0] if context.pages else context.new_page()
        else:
            browser = p.chromium.launch(**launch_kw)
            context = browser.new_context(**context_kw)
            page = context.new_page()

        run_error: BaseException | None = None
        try:
            _run_herts_apply_start(page, xlsx)
        except BaseException as e:
            run_error = e
            print("\n[错误] 脚本执行失败，浏览器将保持打开以便检查页面：")
            traceback.print_exc()
        finally:
            input(">>> 按 Enter 关闭浏览器...")
            _safe_close_browser(context, browser)

        if run_error is not None:
            raise run_error


if __name__ == "__main__":
    main()
